"""指定フォルダ内の.tdfファイルをすべて解析し、抽出結果を1つのExcelにまとめる。

使い方:
    python extract_to_excel.py <tdfフォルダ> [<tdfフォルダ2> ...] <出力xlsxパス>

フォルダ内の*.tdfを再帰的に検索し、ファイルごとに製品マスタ情報
(図番・製品マーク・設計符号・サイズ・本数・重量・長さ・左継手・右継手・
種別(普通梁/斜め梁))を抽出して1シートの表にまとめる。複数フォルダを
指定した場合は指定順にまとめて1つの表に出力する。抽出できなかった項目は
空欄で出力し、処理自体が失敗したファイルはエラー内容を別シート「エラー」
に記録する。
"""
from __future__ import annotations

import math
import re
import sys
from pathlib import Path

import tdf_master_extractor as ex
import tdf_master_extractor_multi as exm
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = [
    "ファイル名", "図番", "製品マーク", "設計符号", "サイズ", "本数",
    "長さ", "重量", "左継手", "右継手", "種別", "製品段",
]

_AXIS_TOLERANCE_DEG = 2.0

# 1枚の図面に梁の立面が2段(以上)に描かれている場合、下段の行から見て
# 上段の行より上は別の梁の寸法線なので長さ判定に含めてはいけない。
# 同一グループ内の行同士のY間隔(実測240前後)と、段違いのY間隔(実測7700
# 以上)の間に十分な余裕を持たせた係数として3000を採用する
# (2026-09-14、EA1-1B-01のB1000A/B90A2グループで確認)。
TIER_Y_GAP_THRESHOLD = 3000.0



def sort_rows(rows: list) -> None:
    """製品の並び順を「上にある製品が先→同じ高さなら左の製品が先→同じ
    製品内はマーク番号の昇順」にする(2026-09-14、EA1-1B-01でB90A[上]→
    B1000A[下]の順になるよう指示を受けて追加)。

    「同じ製品」はマーク末尾の番号を除いた接頭辞が一致するマーク同士と
    みなす(物理的に複数の表・列に分かれて描かれていても、同じ接頭辞なら
    まとめて番号順に並べる。2026-09-14、EA1-1B-06でB34マークが複数の表に
    分かれていた事例で確認)。その製品の代表Y座標は最大Y(最も上にある
    行のY)を使う。"""
    prefix_max_y: dict[str, float] = {}
    prefix_min_x: dict[str, float] = {}
    for r in rows:
        p = exm.mark_prefix(r.mark)
        prefix_max_y[p] = max(prefix_max_y.get(p, r.y), r.y)
        prefix_min_x[p] = min(prefix_min_x.get(p, r.x_mark), r.x_mark)

    def key_func(r):
        p = exm.mark_prefix(r.mark)
        return (-prefix_max_y[p], prefix_min_x[p], exm.mark_sort_key(r.mark))

    rows.sort(key=key_func)


def compute_tier_info(rows: list) -> dict:
    """列(X座標)に関係なく、ファイル全体を対象にY座標で「段」クラスタリング
    する。製品マーク群のY座標が(TIER_Y_GAP_THRESHOLD以上)大きく異なる
    グループを、ページの上半分=1段・下半分=2段とみなす(2026-09-14、
    ユーザー確認: 段はX座標の列ではなくページ全体のY位置で決まる。
    EA1-1B-09で、列が異なるEA12-1B34-13[上]とEA12-1B29-51/52/53[下]が
    同じ2段になる実例で確認)。

    戻り値は id(row) -> (段ラベル, 長さ判定に使うy_max) の辞書。
    Yが最も高い(図面上で最も上にある)段を「1段」とし、以降下に行くほど
    「2段」「3段」…と番号を振る。最上段(1段)はy_max=None(従来通り上方向
    無制限)。それ以外の段は、直上の段の最小Y座標をy_maxとして長さ判定の
    範囲を打ち切る(自分の段の寸法線が、直上の段の寸法線と混ざらない
    ようにするため)。"""
    if not rows:
        return {}

    sorted_rows = sorted(rows, key=lambda r: -r.y)
    tiers = [[sorted_rows[0]]]
    for r in sorted_rows[1:]:
        if tiers[-1][-1].y - r.y > TIER_Y_GAP_THRESHOLD:
            tiers.append([r])
        else:
            tiers[-1].append(r)

    info = {}
    for i, tier in enumerate(tiers):
        tier_label = f"{i + 1}段"
        y_max = None if i == 0 else min(r.y for r in tiers[i - 1])
        for r in tier:
            info[id(r)] = (tier_label, y_max)
    return info


# ---------------------------------------------------------------------------
# 継手判定(大梁専用: 小梁のassign_joints_batchとは別コード)
# ---------------------------------------------------------------------------
#
# 2026-09-14当初は「GJで始まるテキストのみ継手候補」という前提だったが、
# WA2-3G-05(本社/3階、48ファイル基準には含まれない現場・階)で`GW100G`という
# GJ以外の継手コードが使われている実例が見つかった。大梁でも小梁と同じく
# 「長丸で囲まれたテキストは全て継手マークの一種」という前提に立ち、
# GJ限定を撤廃する。
#
# ただし大梁の図面では、長丸が継手コードだけでなく柱マーク(P441等)も
# 囲んでいることがあり、単純に全テキストを候補にすると柱マークを誤って
# 継手として拾ってしまう(2026-09-14、EA1-1G-13/14・EA2-1G-15/16の
# `P441`で確認)。柱マークは、その長丸のすぐ近く(上下左右1000mm以内)に
# 同じテキストが別途(円で囲まれていない形で)描かれている、という特徴が
# ある(実際の柱位置を示すラベルが別にあるため。ユーザー指摘により発見)。
# 一方、継手コード自身は近傍に重複が無い。これを候補除外の条件にする。
#
# 数値のみのテキスト(寸法値等)は比較対象から除外する。取付ピッチ等の
# 寸法値がたまたま継手コードの数字部分と部分一致してしまい、無関係な
# 数値によって誤って除外されることがあったため
# (例: "441"という寸法値が"TB441"に部分一致してしまう小梁側での検証で判明)。
NEARBY_DUP_RANGE = 1000.0


def _has_nearby_duplicate(tdf, mx: float, my: float, text: str) -> bool:
    for rec in tdf.texts:
        if abs(rec.x - mx) < 3 and abs(rec.y - my) < 3:
            continue  # 長丸内の自分自身は除く
        if abs(rec.x - mx) > NEARBY_DUP_RANGE or abs(rec.y - my) > NEARBY_DUP_RANGE:
            continue
        other = tdf.resolve_text(rec)
        if not other:
            continue
        if other.strip().isdigit():
            continue  # 寸法値等の純粋な数値は比較対象外
        if other == text or text in other or other in text:
            return True
    return False


def assign_joints_batch(tdf, rows: list, tier_info: dict, lengths: dict) -> dict:
    """大梁(1G系)専用の継手判定。小梁の`exm.assign_joints_batch`と同じ
    考え方(全長丸テキスト候補+実部材線からの2D距離+ファイル内相互排他)
    だが、柱マーク除外(`_has_nearby_duplicate`)を追加している点が異なる。
    小梁側のコードには一切手を加えず、大梁専用としてこちらに実装する。

    戻り値は id(row) -> (left, right)。"""
    raw_groups: dict[tuple, list] = {}
    for r in rows:
        key = (round(r.x_mark, 1), round(r.x_next, 1))
        raw_groups.setdefault(key, []).append(r)

    groups: dict[tuple, list] = {}
    for key, members in raw_groups.items():
        members_by_y = sorted(members, key=lambda r: r.y)
        clusters: list[list] = [[members_by_y[0]]]
        for r in members_by_y[1:]:
            if r.y - clusters[-1][-1].y > TIER_Y_GAP_THRESHOLD:
                clusters.append([r])
            else:
                clusters[-1].append(r)
        for i, cluster in enumerate(clusters):
            groups[(key, i)] = cluster

    centers = ex.find_stadium_centers(tdf)
    candidates_raw = []
    for mx, my, _r in centers:
        near = [rec for rec in tdf.texts if abs(rec.x - mx) < 3 and abs(rec.y - my) < 3]
        for rec in near:
            t = tdf.resolve_text(rec)
            if not t:
                continue
            if _has_nearby_duplicate(tdf, mx, my, t):
                continue
            candidates_raw.append((mx, my, t))

    claims: dict[tuple, tuple] = {}
    for key, members in groups.items():
        member_lengths = [lengths[id(r)] for r in members if lengths.get(id(r)) is not None]
        if not member_lengths:
            continue
        length_value = max(member_lengths)
        main_axis_deg = 0.0
        theta = -math.radians(main_axis_deg)

        def rotate(x, y, theta=theta):
            xr = x * math.cos(theta) - y * math.sin(theta)
            yr = x * math.sin(theta) + y * math.cos(theta)
            return xr, yr

        threshold = max(length_value * 0.6, exm.JOINT_MIN_THRESHOLD)
        for r in members:
            _tier_label, y_max = tier_info.get(id(r), (None, None))
            row_length = lengths.get(id(r))
            line_endpoints = None
            if row_length is not None:
                line_endpoints = exm._find_reference_line_endpoints(
                    tdf, r.x_mark, r.x_next, r.y, y_max, row_length,
                )
            if line_endpoints is not None:
                (lx, ly), (rx, ry) = line_endpoints
                left_ref, right_ref = (lx, ly), (rx, ry)
                use_2d = True
            else:
                x_mark_r, _ = rotate(r.x_mark, r.y)
                x_next_r, _ = rotate(r.x_next, r.y)
                left_ref, right_ref = (x_mark_r, None), (x_next_r, None)
                use_2d = False
            for mx, my, t in candidates_raw:
                if my <= r.y or (y_max is not None and my >= y_max):
                    continue
                xr, yr = rotate(mx, my)
                for side, ref in (("left", left_ref), ("right", right_ref)):
                    if use_2d:
                        dist = math.hypot(mx - ref[0], my - ref[1])
                    else:
                        dist = abs(xr - ref[0])
                    if dist > threshold:
                        continue
                    cand_key = (round(mx, 2), round(my, 2), t, side)
                    cur = claims.get(cand_key)
                    if cur is None or dist < cur[0]:
                        claims[cand_key] = (dist, key)

    best_per_group_side: dict[tuple, tuple] = {}
    for (_mx, _my, t, side), (dist, key) in claims.items():
        cur = best_per_group_side.get((key, side))
        if cur is None or dist < cur[0]:
            best_per_group_side[(key, side)] = (dist, t)

    group_result: dict[tuple, list] = {}
    for (key, side), (_dist, t) in best_per_group_side.items():
        left_right = group_result.setdefault(key, [None, None])
        left_right[0 if side == "left" else 1] = t

    row_result = {}
    for key, members in groups.items():
        left, right = group_result.get(key, (None, None))
        for r in members:
            row_result[id(r)] = (left, right)
    return row_result


def classify_beam_type(main_axis_deg: float | None) -> str | None:
    """主軸方向(main_axis_deg)から「普通梁」(0度/90度=図面の縦横に沿う
    向き)か「斜め梁」(それ以外の角度)かを判定する。長さが判定できず
    main_axis_degが無い場合はNoneを返す。"""
    if main_axis_deg is None:
        return None
    remainder = abs(main_axis_deg) % 90.0
    if remainder <= _AXIS_TOLERANCE_DEG or remainder >= 90.0 - _AXIS_TOLERANCE_DEG:
        return "普通梁"
    return "斜め梁"

_SEQ_NUMBER_PATTERN = re.compile(r"-(\d+)([A-Za-z]*)　")


def _natural_sort_key(name: str) -> tuple:
    """ファイル名の連番部分(例: "-05"や"-05A")を数値として比較し、
    英字無し("-05")が英字付き("-05A")より先に来るようにする自然順ソート
    キー(2026-09-14、WA2-1G-05/WA2-1G-05Aの並び順修正で追加)。"""
    m = _SEQ_NUMBER_PATTERN.search(name)
    if not m:
        return (name,)
    prefix = name[:m.start()]
    number = int(m.group(1))
    suffix = m.group(2)
    rest = name[m.end():]
    return (prefix, number, suffix, rest)


def collect_tdf_files(folder: Path) -> list[Path]:
    return sorted(folder.rglob("*.tdf"), key=lambda p: _natural_sort_key(p.name))


def main() -> None:
    if len(sys.argv) < 3:
        print("使い方: python extract_to_excel.py <tdfフォルダ> [<tdfフォルダ2> ...] <出力xlsxパス>")
        raise SystemExit(1)

    folders = [Path(p) for p in sys.argv[1:-1]]
    out_path = Path(sys.argv[-1])

    tdf_files: list[Path] = []
    for folder in folders:
        if not folder.exists():
            print(f"フォルダが見つかりません: {folder}")
            raise SystemExit(1)
        tdf_files.extend(collect_tdf_files(folder))
    print(f"対象ファイル数: {len(tdf_files)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "抽出結果"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    errors = []
    total_rows = 0
    for tdf_path in tdf_files:
        try:
            tdf = ex.tb.load(str(tdf_path))
            drawing_number = ex.find_drawing_number(tdf)
            rows = ex.find_product_rows(tdf)
            # 「1製品に複数マークが共有設計符号・サイズで並ぶ」パターン(小梁等)は
            # 別コード(tdf_master_extractor_multi)で検出し、既存ロジックが拾った
            # daiセルとは重複しないようにマージする(2026-09-14追加)。
            existing_dai_positions = {(round(r.x_next, 1), round(r.y, 1)) for r in rows}
            rows = rows + exm.find_product_rows_shared_group(tdf, existing_positions=existing_dai_positions)
            sort_rows(rows)
            tier_info = compute_tier_info(rows)

            lengths: dict[int, float | None] = {}
            beam_types: dict[int, str | None] = {}
            for row in rows:
                tier_label, tier_y_max = tier_info.get(id(row), (None, None))
                if len(rows) == 1:
                    x_min = row.x_mark - ex.RELAX_MARGIN
                    x_max = row.x_next + ex.RELAX_MARGIN
                else:
                    x_min = row.x_mark
                    x_max = row.x_next
                length_value, debug = ex.determine_length(tdf, x_min, x_max, row.y, y_max=tier_y_max)
                lengths[id(row)] = length_value
                beam_types[id(row)] = classify_beam_type(
                    debug.get("main_axis_deg") if length_value is not None else None
                )

            joints = assign_joints_batch(tdf, rows, tier_info, lengths)

            for row in rows:
                tier_label, _tier_y_max = tier_info.get(id(row), (None, None))
                length_value = lengths[id(row)]
                left, right = joints.get(id(row), (None, None))
                if length_value is None:
                    left = right = None
                ws.append([
                    tdf_path.name, drawing_number, row.mark, row.design_code,
                    row.size, row.count, length_value, row.weight, left, right,
                    beam_types[id(row)], tier_label,
                ])
                total_rows += 1
            if not rows:
                errors.append((tdf_path.name, "製品情報テーブルが見つかりませんでした"))
        except Exception as e:  # noqa: BLE001 一括処理のため個別ファイルの失敗で全体を止めない
            errors.append((tdf_path.name, f"{type(e).__name__}: {e}"))

    for col_idx, header in enumerate(HEADERS, start=1):
        width = max(12, len(header) * 2)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    if errors:
        ws_err = wb.create_sheet("エラー")
        ws_err.append(["ファイル名", "内容"])
        for cell in ws_err[1]:
            cell.font = Font(bold=True)
        for name, msg in errors:
            ws_err.append([name, msg])
        ws_err.column_dimensions["A"].width = 40
        ws_err.column_dimensions["B"].width = 60

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"完了: {total_rows}行を抽出しました({len(errors)}件のエラー)")
    print(f"出力先: {out_path}")


if __name__ == "__main__":
    main()
